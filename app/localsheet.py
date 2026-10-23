"""Read-only local spreadsheet source with the same read API the dashboard uses from gspread.

LOCAL_DATA_PATH is either
  * an Excel workbook (.xlsx / .xlsm) whose tabs are named Leads, Referrals, Outreach_Log, Campaign_Queue, or
  * a folder holding Leads.csv, Referrals.csv, Outreach_Log.csv, Campaign_Queue.csv.
Files are never written to.
"""
import csv
import datetime
import os
import shutil
import tempfile


class LocalSheetError(Exception):
    pass


def _s(v):
    if v is None:
        return ""
    if isinstance(v, datetime.datetime):
        return v.date().isoformat() if v.time() == datetime.time(0) else v.isoformat(sep=" ")
    if isinstance(v, datetime.date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


class _Tab:
    def __init__(self, rows):
        self.rows = rows

    def get_all_values(self):
        return [list(r) for r in self.rows]


class LocalSheet:
    def __init__(self, path):
        self.path = path
        if not os.path.exists(path):
            raise LocalSheetError("LOCAL_DATA_PATH does not exist")
        self.is_dir = os.path.isdir(path)

    def worksheet(self, name):
        return _Tab(self._read_dir(name) if self.is_dir else self._read_xlsx(name))

    def _read_dir(self, name):
        for fn in os.listdir(self.path):
            if fn.lower() == f"{name.lower()}.csv":
                with open(os.path.join(self.path, fn), newline="", encoding="utf-8-sig") as f:
                    return [list(r) for r in csv.reader(f)]
        raise LocalSheetError(f"no {name}.csv in the folder")

    def _read_xlsx(self, name):
        import openpyxl
        # Copy first so a workbook that is open in Excel (which locks the file) can still be read.
        with tempfile.TemporaryDirectory() as tmp:
            copy = os.path.join(tmp, "wb.xlsx")
            try:
                shutil.copyfile(self.path, copy)
            except OSError:
                raise LocalSheetError("could not read the workbook (is it locked by another program?)")
            wb = openpyxl.load_workbook(copy, read_only=True, data_only=True)
            try:
                tab = next((n for n in wb.sheetnames if n.lower() == name.lower()), None)
                if tab is None:
                    raise LocalSheetError(f"no '{name}' tab in the workbook")
                rows = [[_s(c) for c in r] for r in wb[tab].iter_rows(values_only=True)]
            finally:
                wb.close()
        while rows and not any(rows[-1]):
            rows.pop()  # trailing blank rows from Excel
        return rows
